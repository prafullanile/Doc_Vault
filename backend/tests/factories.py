"""Builds real test documents in memory: text PDFs, scanned PDFs, encrypted PDFs, DOCX, images."""

import io
import uuid
import zipfile

import docx
import pymupdf
from PIL import Image, ImageDraw, ImageFont

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

SAMPLE_ENGLISH = (
    "The company reported strong results this year. Revenue increased by twenty three percent "
    "compared with the previous year, driven by growth in the services business."
)


def text_pdf(pages: list[str], marker: str | None = None) -> bytes:
    """A PDF with a real text layer. `marker` makes the bytes (and checksum) unique."""
    document = pymupdf.open()
    for text in pages:
        page = document.new_page()
        if text:
            page.insert_textbox(pymupdf.Rect(72, 72, 540, 770), text, fontsize=11)
    document.set_metadata({"title": marker or uuid.uuid4().hex})
    data: bytes = document.tobytes()
    document.close()
    return data


def text_image(text: str, size: tuple[int, int] = (1600, 400)) -> Image.Image:
    image = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.load_default(size=48)
    draw.multiline_text((40, 40), text, fill="black", font=font, spacing=12)
    return image


def png(text: str) -> bytes:
    buf = io.BytesIO()
    text_image(text).save(buf, format="PNG")
    return buf.getvalue()


def scanned_pdf(text: str) -> bytes:
    """A PDF page that is only an image of text — no text layer, so it needs OCR."""
    document = pymupdf.open()
    page = document.new_page()
    page.insert_image(pymupdf.Rect(36, 36, 576, 171), stream=png(text))
    data: bytes = document.tobytes()
    document.close()
    return data


def encrypted_pdf() -> bytes:
    document = pymupdf.open()
    document.new_page().insert_text((72, 72), "secret contents")
    data: bytes = document.tobytes(
        encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="pw", owner_pw="owner"
    )
    document.close()
    return data


def docx_file(sections: list[tuple[str, str]], table: list[list[str]] | None = None) -> bytes:
    document = docx.Document()
    for heading, body in sections:
        document.add_heading(heading, level=1)
        document.add_paragraph(body)
    if table:
        t = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def zip_bomb_docx(expanded_bytes: int) -> bytes:
    """A DOCX-shaped archive with a highly compressible entry of `expanded_bytes`."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", b"\0" * expanded_bytes)
    return buf.getvalue()
